"""Build the property development feasibility workbook.

Left half is the deal you type in. Right half is a dashboard that answers the
only question that matters: does this clear the profit threshold.

    python build_workbook.py

Writes output/property_feasibility.xlsx and output/cellmap.json
"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation

# --- design tokens ---------------------------------------------------------

NAVY = "1F3A5F"
NAVY_DK = "16293F"
INK = "1A1A1A"
MUTED = "7A7A7A"
LINE = "D9DEE5"

INPUT_FILL = PatternFill("solid", fgColor="E8F1FB")
TILE_FILL = PatternFill("solid", fgColor="F4F6F9")
PANEL_FILL = PatternFill("solid", fgColor="FFFFFF")
BAND_FILL = PatternFill("solid", fgColor=NAVY)
SUBBAND_FILL = PatternFill("solid", fgColor="E4E9F0")

# Differential fills (conditional formatting) take the colour in bgColor.
GOOD_CF = PatternFill(bgColor="CDE9D6")
BAD_CF = PatternFill(bgColor="F6D5D2")

HAIR = Side(style="thin", color=LINE)
BOX = Border(left=HAIR, right=HAIR, top=HAIR, bottom=HAIR)
UNDER = Border(bottom=HAIR)

MONEY = '#,##0;[Red]-#,##0'
MONEY0 = '$#,##0;[Red]-$#,##0'
PCT = '0.0%'
PCT_S = '0%'

DEFAULTS = {
    "Residential": {
        "land_comps": [292000, 305000, 298000, 305000],
        "retail_comps": [955000, 972000, 960000, 973000],
        "units": 6, "build_cost": 310000, "council": 28000,
        "stamp_duty_pct": 0.055, "prof_fees_pct": 0.07, "contingency_pct": 0.05,
        "selling_pct": 0.025, "finance_pct": 0.06, "threshold": 0.25,
    },
    "Townhouse": {
        "land_comps": [442000, 455000, 448000, 455000],
        "retail_comps": [1160000, 1178000, 1165000, 1181000],
        "units": 8, "build_cost": 395000, "council": 34000,
        "stamp_duty_pct": 0.055, "prof_fees_pct": 0.08, "contingency_pct": 0.06,
        "selling_pct": 0.025, "finance_pct": 0.065, "threshold": 0.25,
    },
}

SHIFTS = [-0.10, -0.05, 0.0, 0.05, 0.10]


# --- small helpers ---------------------------------------------------------


def band(ws, row, c0, c1, text, *, size=10, fill=BAND_FILL, colour="FFFFFF", height=20):
    ws.merge_cells(start_row=row, start_column=c0, end_row=row, end_column=c1)
    for c in range(c0, c1 + 1):
        ws.cell(row=row, column=c).fill = fill
    cell = ws.cell(row=row, column=c0, value=text)
    cell.font = Font(bold=True, size=size, color=colour)
    cell.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[row].height = height


def text(ws, row, col, value, *, size=9.5, bold=False, colour=INK, italic=False,
         align=None, indent=0, wrap=False):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(size=size, bold=bold, color=colour, italic=italic)
    c.alignment = Alignment(horizontal=align, vertical="center",
                            indent=indent, wrap_text=wrap)
    return c


def input_cell(ws, row, col, value, fmt=MONEY):
    c = ws.cell(row=row, column=col, value=value)
    c.fill = INPUT_FILL
    c.border = BOX
    c.number_format = fmt
    c.font = Font(size=9.5, color=INK)
    c.alignment = Alignment(horizontal="right", indent=1)
    c.protection = Protection(locked=False, hidden=False)
    return f"{c.column_letter}{row}"


def calc_cell(ws, row, col, formula, fmt=MONEY, *, bold=False, size=9.5,
              fill=None, colour=INK):
    c = ws.cell(row=row, column=col, value=formula)
    c.border = BOX
    c.number_format = fmt
    c.font = Font(size=size, bold=bold, color=colour)
    c.alignment = Alignment(horizontal="right", indent=1)
    c.protection = Protection(locked=True, hidden=True)
    if fill:
        c.fill = fill
    return f"{c.column_letter}{row}"


def anchor(ref: str) -> str:
    col = "".join(ch for ch in ref if ch.isalpha())
    row = "".join(ch for ch in ref if ch.isdigit())
    return f"${col}${row}"


def tile(ws, r0, c0, c1, label_text, formula, fmt, *, big=18, colour=NAVY):
    """A KPI card: small muted label over a large number."""
    ws.merge_cells(start_row=r0, start_column=c0, end_row=r0, end_column=c1)
    ws.merge_cells(start_row=r0 + 1, start_column=c0, end_row=r0 + 2, end_column=c1)
    for rr in range(r0, r0 + 3):
        for cc in range(c0, c1 + 1):
            cell = ws.cell(row=rr, column=cc)
            cell.fill = TILE_FILL
            cell.border = BOX
    lab = ws.cell(row=r0, column=c0, value=label_text.upper())
    lab.font = Font(bold=True, size=8, color=MUTED)
    lab.alignment = Alignment(horizontal="center", vertical="center")
    val = ws.cell(row=r0 + 1, column=c0, value=formula)
    val.font = Font(bold=True, size=big, color=colour)
    val.number_format = fmt
    val.alignment = Alignment(horizontal="center", vertical="center")
    val.protection = Protection(locked=True, hidden=True)
    # Deliberately no row heights here: the left-hand input column shares these
    # rows, and setting heights per tile knocks its labels out of line.
    return f"{val.column_letter}{r0 + 1}"


# --- the deal sheet --------------------------------------------------------


def build_deal_sheet(ws, name: str, d: dict) -> dict:
    ws.sheet_view.showGridLines = False
    widths = {"A": 33, "B": 15, "C": 2.5, "D": 13.5, "E": 13.5, "F": 13.5,
              "G": 13.5, "H": 13.5, "I": 13.5}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    # --- masthead ---
    band(ws, 1, 1, 9, f"  {name.upper()} DEVELOPMENT  ·  FEASIBILITY", size=13, height=30)
    text(ws, 2, 1, "Fill the blue cells. Everything else is calculated and locked.",
         size=9, italic=True, colour=MUTED, indent=1)
    ws.row_dimensions[2].height = 16

    # ================= LEFT: inputs =================
    r = 4
    band(ws, r, 1, 2, "THE DEAL", size=9, fill=SUBBAND_FILL, colour=NAVY, height=18)
    r += 1

    text(ws, r, 1, "Raw land comparables", size=9, bold=True, colour=NAVY)
    r += 1
    land_rows = []
    for i in range(4):
        text(ws, r, 1, f"Comparable {i + 1}", size=9, colour=MUTED, indent=1)
        land_rows.append(input_cell(ws, r, 2, d["land_comps"][i]))
        r += 1

    text(ws, r, 1, "Retail comparables", size=9, bold=True, colour=NAVY)
    r += 1
    retail_rows = []
    for i in range(4):
        text(ws, r, 1, f"Comparable {i + 1}", size=9, colour=MUTED, indent=1)
        retail_rows.append(input_cell(ws, r, 2, d["retail_comps"][i]))
        r += 1

    text(ws, r, 1, "Number of dwellings", size=9)
    units = input_cell(ws, r, 2, d["units"], fmt="0")
    r += 1
    text(ws, r, 1, "Build cost per dwelling", size=9)
    build = input_cell(ws, r, 2, d["build_cost"])
    r += 1
    text(ws, r, 1, "Council contributions per dwelling", size=9)
    council = input_cell(ws, r, 2, d["council"])
    r += 2

    band(ws, r, 1, 2, "RATES", size=9, fill=SUBBAND_FILL, colour=NAVY, height=18)
    r += 1
    rate_defs = [
        ("Stamp duty", "stamp_duty_pct"), ("Professional fees", "prof_fees_pct"),
        ("Contingency", "contingency_pct"), ("Selling costs", "selling_pct"),
        ("Finance costs", "finance_pct"),
    ]
    rates = {}
    for lbl, key in rate_defs:
        text(ws, r, 1, lbl, size=9)
        rates[key] = input_cell(ws, r, 2, d[key], fmt=PCT)
        r += 1
    stamp, prof = rates["stamp_duty_pct"], rates["prof_fees_pct"]
    cont, sell = rates["contingency_pct"], rates["selling_pct"]
    fin = rates["finance_pct"]

    r += 1
    band(ws, r, 1, 2, "THE TEST", size=9, fill=SUBBAND_FILL, colour=NAVY, height=18)
    r += 1
    text(ws, r, 1, "Profit-on-cost threshold", size=9, bold=True)
    threshold = input_cell(ws, r, 2, d["threshold"], fmt=PCT)
    r += 1
    text(ws, r, 1, "Industry benchmark is 25%", size=8, italic=True, colour=MUTED)
    left_end = r

    # --- hidden working (column K), keeps the visible sheet clean ---
    k = 4
    def work(formula, fmt=MONEY):
        nonlocal k
        ref = calc_cell(ws, k, 11, formula, fmt)
        k += 1
        return ref

    avg_land = work(f"=IFERROR(AVERAGE({land_rows[0]}:{land_rows[-1]}),0)")
    avg_retail = work(f"=IFERROR(AVERAGE({retail_rows[0]}:{retail_rows[-1]}),0)")
    land_purchase = work(f"={avg_land}*{units}")
    stamp_amt = work(f"={land_purchase}*{stamp}")
    total_purchase = work(f"={land_purchase}+{stamp_amt}")
    gross = work(f"={avg_retail}*{units}")
    sell_amt = work(f"={gross}*{sell}")
    net_real = work(f"={gross}-{sell_amt}")
    build_total = work(f"={build}*{units}")
    council_total = work(f"={council}*{units}")
    prof_amt = work(f"={build_total}*{prof}")
    cont_amt = work(f"={build_total}*{cont}")
    fin_amt = work(
        f"=({total_purchase}+{build_total}+{council_total}+{prof_amt}+{cont_amt})*{fin}")
    tdc = work(f"={build_total}+{council_total}+{prof_amt}+{cont_amt}+{fin_amt}")
    total_cost = work(f"={total_purchase}+{tdc}")
    net_profit = work(f"={net_real}-{total_cost}")
    poc = work(f"=IFERROR({net_profit}/{total_cost},0)", PCT)
    ws.column_dimensions["K"].hidden = True

    # ================= RIGHT: dashboard =================
    band(ws, 4, 4, 9, "RESULT", size=9, fill=SUBBAND_FILL, colour=NAVY, height=18)

    t_profit = tile(ws, 6, 4, 6, "Net profit", f"={net_profit}", MONEY0)
    t_poc = tile(ws, 6, 7, 9, "Profit on cost", f"={poc}", PCT)
    t_cost = tile(ws, 10, 4, 6, "Total cost", f"={total_cost}", MONEY0, big=13)
    t_gdv = tile(ws, 10, 7, 9, "Gross realisation", f"={gross}", MONEY0, big=13)

    # --- verdict ---
    v0 = 14
    ws.merge_cells(start_row=v0, start_column=4, end_row=v0 + 2, end_column=9)
    for rr in range(v0, v0 + 3):
        for cc in range(4, 10):
            ws.cell(row=rr, column=cc).border = BOX
        ws.row_dimensions[rr].height = 19
    v = ws.cell(row=v0, column=4)
    v.value = (f'=IF({total_cost}=0,"ENTER YOUR DEAL",'
               f'IF({poc}>={threshold},"MEETS THRESHOLD","BELOW THRESHOLD"))')
    v.font = Font(bold=True, size=20, color=INK)
    v.alignment = Alignment(horizontal="center", vertical="center")
    v.protection = Protection(locked=True, hidden=True)

    a_cost, a_poc, a_thr = anchor(total_cost), anchor(poc), anchor(threshold)
    vrng = f"D{v0}:I{v0 + 2}"
    ws.conditional_formatting.add(
        vrng, FormulaRule(formula=[f"AND({a_cost}<>0,{a_poc}>={a_thr})"], fill=GOOD_CF))
    ws.conditional_formatting.add(
        vrng, FormulaRule(formula=[f"AND({a_cost}<>0,{a_poc}<{a_thr})"], fill=BAD_CF))

    text(ws, v0 + 3, 4, "Verdict tests profit on cost against your threshold.",
         size=8, italic=True, colour=MUTED)

    # --- where the money goes ---
    b0 = v0 + 5
    band(ws, b0, 4, 9, "WHERE THE MONEY GOES", size=9, fill=SUBBAND_FILL,
         colour=NAVY, height=18)
    breakdown = [
        ("Land and stamp duty", total_purchase),
        ("Build", build_total),
        ("Council contributions", council_total),
        ("Professional fees", prof_amt),
        ("Contingency", cont_amt),
        ("Finance", fin_amt),
    ]
    br = b0 + 1
    for lbl, ref in breakdown:
        text(ws, br, 4, lbl, size=9, indent=1)
        ws.merge_cells(start_row=br, start_column=4, end_row=br, end_column=5)
        calc_cell(ws, br, 6, f"={ref}", MONEY)
        calc_cell(ws, br, 7, f"=IFERROR({ref}/{total_cost},0)", PCT_S)
        # In-cell bar: proportion of total cost, drawn across two columns.
        calc_cell(ws, br, 8, f"=IFERROR({ref}/{total_cost},0)", PCT_S,
                  colour="FFFFFF")
        ws.merge_cells(start_row=br, start_column=8, end_row=br, end_column=9)
        ws.row_dimensions[br].height = 15
        br += 1
    ws.conditional_formatting.add(
        f"H{b0 + 1}:H{br - 1}",
        DataBarRule(start_type="num", start_value=0, end_type="num", end_value=0.7,
                    color=NAVY, showValue=True, minLength=0, maxLength=100))
    text(ws, br, 4, "Share of total cost", size=8, italic=True, colour=MUTED)

    # --- sensitivity grid ---
    s0 = br + 2
    band(ws, s0, 4, 9, "SENSITIVITY  ·  PROFIT ON COST", size=9,
         fill=SUBBAND_FILL, colour=NAVY, height=18)
    text(ws, s0 + 1, 4, "Retail value shift  ↓     Build cost shift  →",
         size=8, italic=True, colour=MUTED, indent=1)
    ws.merge_cells(start_row=s0 + 1, start_column=4, end_row=s0 + 1, end_column=9)

    hdr = s0 + 2
    for j, shift in enumerate(SHIFTS):
        c = ws.cell(row=hdr, column=5 + j, value=shift)
        c.number_format = '+0%;-0%;0%'
        c.font = Font(bold=True, size=8.5, color=NAVY)
        c.alignment = Alignment(horizontal="center")
        c.fill = SUBBAND_FILL
        c.border = BOX

    bt, tp, ct = anchor(build_total), anchor(total_purchase), anchor(council_total)
    nr, a_prof, a_cont, a_fin = anchor(net_real), anchor(prof), anchor(cont), anchor(fin)

    for i, gshift in enumerate(SHIFTS):
        rr = hdr + 1 + i
        rh = ws.cell(row=rr, column=4, value=gshift)
        rh.number_format = '+0%;-0%;0%'
        rh.font = Font(bold=True, size=8.5, color=NAVY)
        rh.alignment = Alignment(horizontal="center")
        rh.fill = SUBBAND_FILL
        rh.border = BOX
        for j, _ in enumerate(SHIFTS):
            cc = 5 + j
            col_letter = ws.cell(row=hdr, column=cc).column_letter
            b_shift = f"{col_letter}${hdr}"
            bt_s = f"({bt}*(1+{b_shift}))"
            sub = f"({tp}+{bt_s}+{ct}+{bt_s}*{a_prof}+{bt_s}*{a_cont})"
            cost = f"({sub}+{sub}*{a_fin}-{tp}*{a_fin}+{tp}*{a_fin})"
            cost = f"({sub}*(1+{a_fin}))"
            formula = f"=IFERROR({nr}*(1+$D{rr})/{cost}-1,0)"
            calc_cell(ws, rr, cc, formula, PCT, size=9)

    grid = f"E{hdr + 1}:I{hdr + len(SHIFTS)}"
    ws.conditional_formatting.add(grid, ColorScaleRule(
        start_type="num", start_value=0, start_color="F2A9A2",
        mid_type="num", mid_value=0.25, mid_color="FFF2CC",
        end_type="num", end_value=0.45, end_color="BFE3CA"))

    tail = hdr + len(SHIFTS) + 1
    text(ws, tail, 4, "A 5% fall in retail value hurts about twice as much as a "
                      "5% rise in build cost.", size=8, italic=True, colour=MUTED)
    ws.merge_cells(start_row=tail, start_column=4, end_row=tail, end_column=9)

    # --- validation ---
    dv = DataValidation(type="decimal", operator="between", formula1=0, formula2=1,
                        allow_blank=False, showErrorMessage=True)
    dv.error = "Enter a percentage between 0% and 100%."
    dv.errorTitle = "Out of range"
    ws.add_data_validation(dv)
    for ref in (stamp, prof, cont, sell, fin, threshold):
        dv.add(ws[ref])

    udv = DataValidation(type="whole", operator="greaterThan", formula1=0,
                         allow_blank=False, showErrorMessage=True)
    udv.error = "Number of dwellings must be a whole number above zero."
    udv.errorTitle = "Invalid count"
    ws.add_data_validation(udv)
    udv.add(ws[units])

    # --- protection and print ---
    ws.protection.sheet = True
    ws.protection.enable()
    ws.protection.selectLockedCells = False

    last = max(left_end, tail) + 1

    # One rhythm down the page so the two columns line up with each other.
    for rr in range(4, last + 1):
        if ws.row_dimensions[rr].height in (None, 15):
            ws.row_dimensions[rr].height = 16.5

    ws.print_area = f"A1:I{last}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.page_margins.top = ws.page_margins.bottom = 0.4

    return {
        "land": land_rows, "retail": retail_rows, "units": units, "build": build,
        "council": council, "stamp": stamp, "prof": prof, "cont": cont,
        "sell": sell, "fin": fin, "threshold": threshold,
        "avg_land": avg_land, "avg_retail": avg_retail,
        "land_purchase": land_purchase, "stamp_amt": stamp_amt,
        "total_purchase": total_purchase, "gross": gross, "sell_amt": sell_amt,
        "net_real": net_real, "build_total": build_total,
        "council_total": council_total, "prof_amt": prof_amt,
        "cont_amt": cont_amt, "fin_amt": fin_amt, "tdc": tdc,
        "total_cost": total_cost, "net_profit": net_profit, "poc": poc,
        "verdict": f"D{v0}",
        "tiles": {"profit": t_profit, "poc": t_poc, "cost": t_cost, "gdv": t_gdv},
        "sensitivity": {"top_left": f"E{hdr + 1}", "shifts": SHIFTS},
    }


# --- start here ------------------------------------------------------------


def build_start_here(ws, names: list[str]) -> None:
    ws.sheet_view.showGridLines = False
    for col, w in {"A": 3, "B": 30, "C": 62, "D": 3}.items():
        ws.column_dimensions[col].width = w

    band(ws, 2, 2, 3, "  PROPERTY FEASIBILITY CALCULATOR", size=15, height=34)
    text(ws, 3, 2, "  Two minutes to a yes or no on any deal.", size=10,
         italic=True, colour=MUTED)
    ws.row_dimensions[3].height = 18

    def head(row, t):
        text(ws, row, 2, t, size=11, bold=True, colour=NAVY)
        for c in (2, 3):
            ws.cell(row=row, column=c).border = UNDER
        ws.row_dimensions[row].height = 20

    def item(row, left, right, *, swatch=None):
        c = text(ws, row, 2, left, size=9.5, bold=True, indent=1)
        if swatch:
            c.fill = swatch
        text(ws, row, 3, right, size=9.5, colour=INK, wrap=True)
        ws.row_dimensions[row].height = 26 if len(right) > 72 else 17

    r = 5
    head(r, "How to use it"); r += 1
    for n, line in enumerate([
        f"Open the tab for your deal type — {' or '.join(names)}.",
        "Fill only the blue cells. Everything else calculates and is locked.",
        "Read the verdict. Green clears your profit threshold, red does not.",
        "The threshold is itself an input. Change it and the verdict follows.",
    ], start=1):
        item(r, f"Step {n}", line); r += 1

    r += 1
    head(r, "Reading the sheet"); r += 1
    item(r, "Blue cell", "You type here.", swatch=INPUT_FILL); r += 1
    item(r, "Grey tile", "A headline number, calculated for you."); r += 1
    item(r, "Sensitivity grid", "Profit on cost if retail values and build costs "
                                "move against you. Red is trouble."); r += 1

    r += 1
    head(r, "What it works out"); r += 1
    for left, right in [
        ("Market averages", "Average raw land and retail value from your comparables."),
        ("Total purchase", "Land purchase plus stamp duty."),
        ("Gross realisation", "Average retail value times the number of dwellings."),
        ("Net realisation", "Gross realisation less selling costs."),
        ("Development costs", "Build, council contributions, professional fees, "
                              "contingency and finance."),
        ("Total cost", "Total purchase plus total development costs."),
        ("Net profit", "Net realisation less total cost."),
        ("Profit on cost", "Net profit divided by total cost. The verdict tests this."),
    ]:
        item(r, left, right); r += 1

    r += 1
    head(r, "Before you trust it"); r += 1
    item(r, "It is a screen", "This sizes a deal quickly. It is not a substitute "
                              "for a full feasibility, a valuation or advice."); r += 1
    item(r, "No GST", "Sale prices and costs are treated consistently. GST and the "
                      "margin scheme are not modelled."); r += 1
    item(r, "No timeline", "Finance is a flat percentage, not a drawn-down "
                           "facility over a project duration."); r += 1
    item(r, "Worked example", "Every tab opens filled in. Type over it."); r += 1

    ws.protection.sheet = True
    ws.protection.enable()
    ws.print_area = f"A1:D{r}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def main() -> None:
    wb = Workbook()
    start = wb.active
    start.title = "Start Here"

    refs = {}
    for name, defaults in DEFAULTS.items():
        refs[name] = build_deal_sheet(wb.create_sheet(name), name, defaults)
    build_start_here(start, list(DEFAULTS))

    out = Path(__file__).parent / "output" / "property_feasibility.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)

    cellmap = out.parent / "cellmap.json"
    cellmap.write_text(json.dumps(refs, indent=1), encoding="utf-8")

    print(f"Wrote {out}")
    print(f"Wrote {cellmap}")
    for name, r in refs.items():
        print(f"  {name:12s} poc {r['poc']}  verdict {r['verdict']}")


if __name__ == "__main__":
    main()
